"""Discovery kit — enquiry → discovery-questions document (MD, PDF, DOCX) per segment and service, ready to send the same day."""
from __future__ import annotations
import uuid
from datetime import date
from pathlib import Path
import yaml
from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from app.core import report
from app.core.config import settings

router = APIRouter()
KIT = Path(__file__).with_name("kit.yaml")
OUT = settings.outputs_dir / "discovery"
def kit() -> dict: return yaml.safe_load(KIT.read_text())

class Enquiry(BaseModel):
    client: str
    segment: str = Field("corporate", description="corporate | business | university | college | domain")
    service: str = Field("workshop", description="workshop | training | audit | consulting")
    via_partner: str | None = Field(None, description="e.g. IndusGuru — adds a partner-only commercial section")
    contact_name: str | None = None
    notes: str | None = Field(None, description="anything already known from the enquiry")

@router.get("/kit")
def get_kit(): return kit()

def assemble(e: Enquiry) -> dict:
    k = kit(); seg = k["segments"].get(e.segment); svc = k["services"].get(e.service)
    if not seg or not svc: raise HTTPException(400, f"segment in {list(k['segments'])}, service in {list(k['services'])}")
    sections = [{"letter": chr(65 + i), **s} for i, s in enumerate([seg["section"]] + svc["sections"])]
    if e.via_partner:
        sections.append({"letter": chr(65 + len(sections)), "title": "Commercial and contracting (partner only)", "questions": [
            {"q": f"Does {e.via_partner} contract with the client and sub-contract to us, or do we invoice the client directly?", "why": "Determines margin structure and who we present the number to"},
            {"q": "Is the fee presented as a day rate, a package, or a per-participant price?", "why": "Same total, different optics"},
            {"q": "What margin or referral structure applies?", "why": "Built into the presented number, not added after"}]})
    return {"discovery_id": uuid.uuid4().hex[:8], "client": e.client, "segment": seg, "service": svc, "company": k["company"],
            "sections": sections, "confirms": seg["confirms"] + k["company"]["standard_terms"][:4], "next_steps": k["next_steps"],
            "cannot_confirm": "the fee, the dates, and whether labs run in the client's environment or ours", "notes": e.notes, "via_partner": e.via_partner}

def to_sections(d: dict) -> tuple[str, str, list[dict]]:
    co = d["company"]; title = f"{d['service']['title']} — Discovery Questions for {d['client']}"
    sub = f"{co['name']} · {date.today().strftime('%d %B %Y')}" + (f" · via {d['via_partner']}" if d["via_partner"] else "")
    S = [report.section("Purpose", report.para(f"{co['name']} is interested in delivering this engagement and can confirm capability and availability now. "
            "This note sets out what we need from the client before we put a firm number against it. We have deliberately not attached commercials: "
            "the cost is driven less by delivery days than by scope, environment and participant numbers. A short round of questions now avoids a revised number later."),
         *( [report.para(f"<b>Already known from the enquiry:</b> {d['notes']}")] if d["notes"] else [])),
         report.section(f"About {co['name']}", report.para(co["about"]), report.para(co["availability_note"])),
         report.section("What we can confirm today", report.bullets(d["confirms"]), report.para(f"<b>What we cannot confirm without answers:</b> {d['cannot_confirm']}."))]
    for s in d["sections"]:
        S.append(report.section(f"{s['letter']}. {s['title']}", report.table(["#", "Question", "Why it matters"],
                 [[f"{s['letter']}{i}", q["q"], q["why"]] for i, q in enumerate(s["questions"], 1)], [10, 88, 80])))
    S.append(report.section("How the commercials will be built", report.para(f"<b>Fixed component:</b> {d['service']['fixed']}. Stable; moves only with batch or scope count."),
             report.para("<b>Variable components</b> — each moves the number and none can be estimated responsibly without input:"), report.bullets(d["service"]["variables"]),
             report.para("<b>Reductions available:</b> multiple batches booked together, consecutive scheduling, client-provided environments.")))
    S.append(report.section("Next steps", report.bullets(d["next_steps"])))
    S.append(report.section("Contact", report.para(f"{co['name']} · GSTIN {co['gstin']}<br/>{co['address']}<br/>{co['phone']} · {co['email']} · {co['website']}")))
    return title, sub, S

def to_docx(title: str, sub: str, S: list[dict], path: Path):
    from docx import Document
    from docx.shared import Pt
    doc = Document(); doc.add_heading(title, 0); doc.add_paragraph(sub)
    import re
    def clean(t): return re.sub(r"<[^>]+>", "", t.replace("<br/>", "\n"))
    for s in S:
        doc.add_heading(s["heading"], level=1)
        for b in s["blocks"]:
            if b["type"] == "p": doc.add_paragraph(clean(b["text"]))
            elif b["type"] == "ul":
                for i in b["items"]: doc.add_paragraph(clean(i), style="List Bullet")
            elif b["type"] == "table":
                t = doc.add_table(rows=1, cols=len(b["headers"])); t.style = "Table Grid"
                for i, h in enumerate(b["headers"]): t.rows[0].cells[i].text = h
                for r in b["rows"]:
                    cells = t.add_row().cells
                    for i, c in enumerate(r): cells[i].text = clean(str(c))
    doc.save(path)

@router.post("/generate")
def generate(e: Enquiry):
    d = assemble(e); title, sub, S = to_sections(d); OUT.mkdir(parents=True, exist_ok=True)
    stem = f"discovery-{d['discovery_id']}"
    report.write_both(OUT, stem, title, sub, S); to_docx(title, sub, S, OUT / f"{stem}.docx")
    d.update({"title": title, "markdown_url": f"/api/discovery/file/{stem}.md", "pdf_url": f"/api/discovery/file/{stem}.pdf", "docx_url": f"/api/discovery/file/{stem}.docx",
              "question_count": sum(len(s["questions"]) for s in d["sections"])})
    return d

@router.get("/file/{filename}")
def file(filename: str):
    p = OUT / filename
    if not p.exists() or p.suffix not in {".md", ".pdf", ".docx"}: raise HTTPException(404, "not found")
    return FileResponse(p)
